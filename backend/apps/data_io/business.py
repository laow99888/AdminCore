"""显式业务资源；预检不执行业务扩展，实际写入复用生成模块事务。"""
import csv
import hashlib
import io
import json
from itertools import islice
from importlib import import_module
from types import SimpleNamespace

from django.db import transaction
from rest_framework.exceptions import NotFound, PermissionDenied, ValidationError

from apps.core.codegen.integrations import registered_module
from apps.core.modules import module_enabled
from apps.permissions.field_permissions import business_allowed_fields
from apps.permissions.permissions import has_permission_code

MAX_ROWS = 5000
MAX_BYTES = 10 * 1024 * 1024


def resources():
    from apps.core.modules import MODULES
    result = {}
    for declaration in MODULES:
        if declaration.data_io_resources and module_enabled(declaration.key):
            schema, _, _ = registered_module(declaration.key)
            if (schema['structure'] == 'single' and not schema['details'] and not schema.get('approval')
                    or schema['structure'] == 'master_detail' and schema['details']):
                result[declaration.key] = {'label': schema['title'], 'module': declaration.key, 'business': True,
                    'importable': True, 'exportable': True, 'permissions': (declaration.key + ':read',),
                    'fields': [field['name'] for field in schema['fields']],
                    'required': [field['name'] for field in schema['fields'] if field['required'] and field['form']]}
    return result


def resource_schema(resource):
    if resource not in resources():
        raise NotFound()
    return registered_module(resource)


def schema_digest(schema):
    return hashlib.sha256(json.dumps(schema, sort_keys=True, separators=(',', ':')).encode()).hexdigest()


def require(actor, resource, mode=None, *, export=False):
    schema, model, service = resource_schema(resource)
    codes = ['data_io:export' if export else 'data_io:import', resource + ':read']
    if not export:
        codes.append(resource + ':create')
        if mode == 'update':
            codes.append(resource + ':update')
    if (not actor or not actor.is_active or actor.must_change_password
            or not module_enabled('data_io') or not all(has_permission_code(actor, code) for code in codes)):
        raise PermissionDenied('没有该业务资源的导入导出权限')
    from apps.core.tenancy import current_tenant_id, enabled
    if enabled() and current_tenant_id() is not None:
        from apps.core.models import Tenant, TenantMembership
        tenant_id = current_tenant_id()
        if (not Tenant.objects.filter(pk=tenant_id, status=Tenant.STATUS_ACTIVE).exists()
                or not (actor.is_staff or actor.is_superuser) and not TenantMembership.objects.filter(
                    tenant_id=tenant_id, user=actor, is_active=True).exists()):
            raise PermissionDenied('租户已停用或成员授权已撤销')
    return schema, model, service


def fields_for(actor, resource, action):
    fields = business_allowed_fields(actor, resource, action)
    if fields is None:
        raise PermissionDenied('业务资源尚未登记字段策略')
    return fields


def catalog(actor, resource):
    schema, _, _ = require(actor, resource)
    readable = fields_for(actor, resource, 'read')
    writable = fields_for(actor, resource, 'write')
    result = [{'name': field['name'], 'label': field['label'], 'kind': field['kind'],
             'required': field['required'], 'key': field['unique'] and field['kind'] == 'string'}
            for field in schema['fields'] if field['form'] and field['name'] in readable & writable]
    if schema['structure'] == 'master_detail':
        from .business_documents import detail_catalog
        result.extend(detail_catalog(schema, readable, writable))
    return result


def parse_upload(uploaded):
    name = uploaded.name.lower()
    if not name.endswith(('.csv', '.xlsx')):
        raise ValidationError('仅支持 CSV 或 XLSX')
    raw = uploaded.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise ValidationError('导入源文件最多 10 MiB')
    return raw, parse_rows(raw, name)


def parse_rows(raw, name):
    """拒绝重复表头、公式、压缩炸弹、额外工作表；不依赖 pandas 的自动改列名。"""
    try:
        if name.lower().endswith('.csv'):
            rows = list(islice(csv.reader(io.StringIO(raw.decode('utf-8-sig'))), MAX_ROWS + 2))
        else:
            from zipfile import ZipFile
            from openpyxl import load_workbook
            with ZipFile(io.BytesIO(raw)) as archive:
                if len(archive.infolist()) > 1000 or sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError()
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            try:
                if len(book.worksheets) != 1:
                    raise ValueError()
                rows = []
                for row in book.active.iter_rows():
                    if len(rows) > MAX_ROWS or len(row) > 128 or any(cell.data_type == 'f' for cell in row):
                        raise ValueError()
                    rows.append([str(cell.value.isoformat() if hasattr(cell.value, 'isoformat') else cell.value)
                                 if cell.value is not None else '' for cell in row])
            finally:
                book.close()
        if not 2 <= len(rows) <= MAX_ROWS + 1:
            raise ValueError()
        headers = [str(value).strip() for value in rows[0]]
        if (not 1 <= len(headers) <= 128 or any(not value or len(value) > 100 for value in headers)
                or len(set(headers)) != len(headers)):
            raise ValueError()
        result = []
        for row in rows[1:]:
            if len(row) > len(headers) or any(len(str(value)) > 32767 or '\0' in str(value) for value in row):
                raise ValueError()
            result.append(dict(zip(headers, [str(value).strip() for value in row] + [''] * (len(headers) - len(row)))))
        return headers, result
    except (ValidationError, PermissionDenied):
        raise
    except Exception as exc:
        raise ValidationError('文件无效：需唯一表头、单工作表、无公式且最多5000行/128列') from exc


def options_for(actor, resource, value, headers=None):
    if not isinstance(value, dict) or set(value) != {'mapping', 'key', 'mode'}:
        raise ValidationError('导入设置需包含列映射、业务键和重复策略')
    mapping, key, mode = value['mapping'], value['key'], value['mode']
    if mode not in ('create', 'skip', 'update'):
        raise ValidationError('重复策略无效')
    require(actor, resource, mode)
    fields = catalog(actor, resource)
    known = {field['name'] for field in fields}
    if (not isinstance(mapping, dict) or not 1 <= len(mapping) <= 128
            or any(not isinstance(source, str) or not source or len(source) > 100
                   or not isinstance(target, str) or target not in known for source, target in mapping.items())
            or len(set(mapping.values())) != len(mapping) or headers is not None and not set(mapping) <= set(headers)):
        raise ValidationError('列映射无效、重复或字段无权访问')
    if not isinstance(key, str) or key not in {field['name'] for field in fields if field['key']} or key not in mapping.values():
        raise ValidationError('请选择有读取和写入权限的唯一文本业务键')
    schema, _, _ = resource_schema(resource)
    if {field['name'] for field in schema['fields'] if field['required'] and field['form']} - set(mapping.values()):
        raise ValidationError('缺少必填字段映射')
    if schema['structure'] == 'master_detail':
        from .business_documents import validate_mapping
        validate_mapping(schema, mapping)
    return {'mapping': dict(mapping), 'key': key, 'mode': mode}


def serialized_values(actor, schema, source, options, *, version=0, instance=None):
    if schema['structure'] == 'master_detail':
        from .business_documents import document_values
        values = document_values(schema, source, options)
    else:
        fields = {field['name']: field for field in schema['fields']}
        values = {target: cell_value(fields[target], source[column]) for column, target in options['mapping'].items()}
    values['version'] = version
    serializer_class = getattr(import_module('apps.' + schema['module'] + '.serializers'), schema['model'] + 'Serializer')
    serializer = serializer_class(instance=instance, data=values, partial=instance is not None,
                                  context={'request': SimpleNamespace(user=actor)})
    serializer.is_valid(raise_exception=True)
    return serializer.validated_data


def cell_value(field, value):
    # 空值对齐页面；日期/关联/附件为null，其他文本保留空字符串。
    if value == '' and field['kind'] in {'date', 'datetime', 'relation', 'file'}:
        return None
    if field['kind'] == 'dictionary' and value:
        from apps.core.codegen.runtime import dictionary_options
        choices = dictionary_options(field)
        by_label = [item['value'] for item in choices if item['label'] == value]
        if value not in {item['value'] for item in choices}:
            if len(by_label) != 1:
                raise ValidationError('字典标签不唯一或不可用')
            return by_label[0]
    return value


def validate_values(actor, schema, service, values, instance):
    from apps.core.codegen.runtime import enforce_business_write
    enforce_business_write(actor, schema, values)
    service.validate_references(schema['fields'], values, actor, instance)
    for group in schema['details']:
        children = values[group['name']]
        if not group['min_rows'] <= len(children) <= group['max_rows']:
            raise ValidationError('明细行数不符合业务定义')
        for child in children:
            service.validate_references(group['fields'], child, actor)


def import_units(schema, rows, options):
    if schema['structure'] == 'master_detail':
        from .business_documents import group_rows
        return group_rows(schema, rows, options)
    return [{'source_rows': [number], 'source': source, 'error': ''} for number, source in enumerate(rows, 2)]


def preflight(actor, resource, rows, options):
    schema, _, service = require(actor, resource, options['mode'])
    units = import_units(schema, rows, options)
    result, seen = [], set()
    summary = dict(total=len(units), created=0, updated=0, skipped=0, failed=0)
    if schema['structure'] == 'master_detail':
        summary['source_rows'] = len(rows)
    key_column = next(column for column, target in options['mapping'].items() if target == options['key'])
    for unit in units:
        source = unit['source']
        item = {'row': unit['source_rows'][0], 'action': 'failed', 'error': unit['error']}
        if schema['structure'] == 'master_detail':
            item['source_rows'] = unit['source_rows']
        key = source[key_column] if schema['structure'] == 'single' else source[0][key_column]
        try:
            if unit['error']:
                raise ValidationError('整单主字段不一致')
            if not key or key in seen:
                item['error'] = 'empty_business_key' if not key else 'duplicate_in_file'
                raise ValidationError('文件内业务键为空或重复')
            seen.add(key)
            row = service.visible(actor).filter(**{options['key']: key}).first()
            if row is not None and options['mode'] == 'create':
                item['error'] = 'business_key_exists'
                raise ValidationError('业务键已存在')
            if row is not None and options['mode'] == 'skip':
                item['action'] = 'skipped'
            else:
                if schema.get('approval') and row is not None and row.approval_status in ('pending', 'approved'):
                    item['error'] = 'approval_locked'
                    raise ValidationError('审批中或已通过的单据不可修改')
                version = row.version if row else 0
                values = serialized_values(actor, schema, source, options, version=version, instance=row)
                with transaction.atomic():
                    validate_values(actor, schema, service, values, row)
                item.update(action='updated' if row else 'created', identifier=str(row.pk) if row else None, version=version)
        except (ValidationError, PermissionDenied, NotFound):
            item['error'] = item['error'] or 'row_validation_failed'
        summary[item['action']] += 1
        result.append(item)
    return result, summary


def business_summary(batch, total):
    receipts = batch.receipts.all()
    result = {key: receipts.filter(outcome=key).count() for key in ('created', 'updated', 'skipped', 'failed')}
    summary = {**result, 'total': total, 'processed': receipts.count(), 'file_id': str(batch.source_id),
            'errors': [{'row': row.row, 'message': row.error} for row in receipts.filter(outcome='failed').order_by('row')[:100]]}
    if 'source_row_count' in batch.options:
        summary['source_rows'] = batch.options['source_row_count']
    return summary
