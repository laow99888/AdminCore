"""真实生成模块的导入导出验收；由隔离验收器安装后执行。"""
import csv
import io
import json
from datetime import timedelta

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.exceptions import ValidationError
from rest_framework.test import APIClient

from apps.cg_suppliers.models import Supplier
from apps.cg_suppliers import extensions
from apps.core.models import Tenant
from apps.core.tenancy import tenant_context
from apps.data_io.business import parse_rows
from apps.data_io.models import BusinessImportPlan, ImportMapping
from apps.data_io.tasks import run_data_import
from apps.logs.models import OperationLog
from apps.permissions.models import Role, Permission, RolePermission, UserRole
from apps.tasks.models import TaskRecord, ImportExecution, ImportRowReceipt

pytestmark = pytest.mark.django_db
ROOT = '/api/v1/data-io/business/'
OPTIONS = {'mapping': {'编号': 'code', '名称': 'name', '启用': 'active'}, 'key': 'code', 'mode': 'create'}


def upload(text='编号,名称,启用\nA,甲,true\n', name='rows.csv'):
    return SimpleUploadedFile(name, text.encode('utf-8-sig'), content_type='text/csv')


def preview(client, text='编号,名称,启用\nA,甲,true\n', mode='create', options=None):
    result = client.post(ROOT + 'preview/', {'resource': 'cg_suppliers', 'file': upload(text),
        'options': json.dumps(options or {**OPTIONS, 'mode': mode})}, format='multipart')
    assert result.status_code == 200, result.content
    return result.json()['data']


def submit(client, plan):
    result = client.post(ROOT + f"plans/{plan['id']}/submit/")
    assert result.status_code == 200, result.content
    return result.json()['data']


def execute(client, plan):
    task = submit(client, plan); run_data_import(task['id'])
    return TaskRecord.objects.get(pk=task['id'])


def restricted(user, *, fields=None, codes=None):
    role = Role.objects.create(name='IO reader', field_permissions=fields or {})
    UserRole.objects.create(user=user, role=role)
    for code in codes or ['cg_suppliers:read', 'cg_suppliers:create', 'cg_suppliers:update', 'data_io:import', 'data_io:export', 'task:retry']:
        resource, action = code.split(':')
        permission, _ = Permission.objects.get_or_create(code=code, defaults={'resource': resource, 'action': action, 'name': code})
        RolePermission.objects.create(role=role, permission=permission)
    client = APIClient(); client.force_authenticate(user)
    return client, role


def test_preview_does_not_mutate_and_confirm_is_idempotent(admin_client, monkeypatch):
    calls = []
    monkeypatch.setattr(extensions, 'before_save', lambda **kwargs: calls.append(kwargs['values']['code']))
    plan = preview(admin_client, '编号,名称,启用\nA,甲,true\nA,重复,false\nB,乙,false\n')
    assert plan['summary'] == {'total': 3, 'created': 2, 'updated': 0, 'skipped': 0, 'failed': 1}
    assert plan['errors'] == [{'row': 3, 'message': 'duplicate_in_file'}]
    assert Supplier.objects.count() == 0 and calls == []
    first = submit(admin_client, plan); again = submit(admin_client, plan)
    assert first['id'] == again['id'] and TaskRecord.objects.filter(task_type='import').count() == 1
    run_data_import(first['id']); run_data_import(first['id'])
    task = TaskRecord.objects.get(pk=first['id'])
    assert task.status == 'success' and task.result['created'] == 2 and task.result['failed'] == 1
    assert calls == ['A', 'B'] and Supplier.objects.count() == 2
    assert ImportRowReceipt.objects.count() == 3
    assert OperationLog.objects.filter(resource='cg_suppliers', task_id=task.pk).count() == 2
    errors = admin_client.get(ROOT + f"import/{task.pk}/errors/")
    assert errors.status_code == 200, errors.content
    rows = list(csv.reader(io.StringIO(errors.content.decode('utf-8-sig'))))
    assert rows[0] == ['source_row', 'error_code', 'code', 'name', 'active']
    assert rows[1][0] == '3' and rows[1][2:] == ['A', '重复', 'false']


@pytest.mark.parametrize('mode,counts', [('create', (0, 0, 0, 1)), ('skip', (0, 0, 1, 0)), ('update', (0, 1, 0, 0))])
def test_duplicate_strategies(admin_client, admin_user, mode, counts):
    row = Supplier.objects.create(code='A', name='旧值', creator=admin_user)
    task = execute(admin_client, preview(admin_client, mode=mode))
    assert tuple(task.result[key] for key in ('created', 'updated', 'skipped', 'failed')) == counts
    row.refresh_from_db(); assert row.name == ('甲' if mode == 'update' else '旧值')


def test_preview_version_conflict_does_not_overwrite(admin_client, admin_user):
    row = Supplier.objects.create(code='A', name='原值', creator=admin_user)
    plan = preview(admin_client, mode='update')
    Supplier.objects.filter(pk=row.pk).update(name='他人新值', version=row.version + 1)
    task = execute(admin_client, plan)
    assert task.result['failed'] == 1 and task.result['updated'] == 0
    assert task.result['errors'] == [{'row': 2, 'message': 'record_changed'}]
    row.refresh_from_db(); assert row.name == '他人新值'


def test_all_preview_errors_download_and_corrected_reupload(admin_client):
    text = '编号,名称,启用\n' + '\n'.join('A,,true' for _ in range(105))
    plan = preview(admin_client, text)
    assert plan['summary']['failed'] == 105 and len(plan['errors']) == 100
    response = admin_client.get(ROOT + f"plans/{plan['id']}/errors/")
    assert response.status_code == 200, response.content
    rows = list(csv.DictReader(io.StringIO(response.content.decode('utf-8-sig'))))
    assert len(rows) == 105 and rows[-1]['source_row'] == '106'
    assert Supplier.objects.count() == 0 and TaskRecord.objects.count() == 0
    # 导出的业务字段名可自动匹配；元数据列忽略，修正后正常重传。
    options = {'mapping': {'code': 'code', 'name': 'name', 'active': 'active'}, 'key': 'code', 'mode': 'create'}
    corrected = preview(admin_client, 'source_row,error_code,code,name,active\n2,row_validation_failed,A,修正,true\n', options=options)
    assert execute(admin_client, corrected).result['created'] == 1


def test_hook_failure_rolls_back_business_and_receipt_is_failed(admin_client, monkeypatch):
    def fail(**kwargs):
        if kwargs['instance'].code == 'A':
            raise ValidationError('extension refuses row')
    monkeypatch.setattr(extensions, 'after_save', fail)
    task = execute(admin_client, preview(admin_client, '编号,名称,启用\nA,甲,true\nB,乙,false\n'))
    assert task.result['created'] == 1 and task.result['failed'] == 1
    assert list(Supplier.objects.values_list('code', flat=True)) == ['B']


@pytest.mark.parametrize('error_type', [RuntimeError, ValueError])
def test_crash_resume_does_not_repeat_committed_rows(admin_client, monkeypatch, error_type):
    original = extensions.after_save
    calls = []
    def crash(**kwargs):
        code = kwargs['instance'].code; calls.append(code)
        if code == 'B':
            raise error_type('simulated worker crash')
        return original(**kwargs)
    monkeypatch.setattr(extensions, 'after_save', crash)
    task = submit(admin_client, preview(admin_client, '编号,名称,启用\nA,甲,true\nB,乙,false\n'))
    with pytest.raises(error_type):
        run_data_import(task['id'])
    assert list(Supplier.objects.values_list('code', flat=True)) == ['A']
    monkeypatch.setattr(extensions, 'after_save', lambda **kwargs: calls.append(kwargs['instance'].code))
    response = admin_client.post('/api/v1/data-io/import/' + task['id'] + '/')
    assert response.status_code == 202, response.content
    new = response.json()['data']; run_data_import(new['id'])
    result = TaskRecord.objects.get(pk=new['id'])
    assert result.result['created'] == 2 and result.result['processed'] == 2
    assert calls == ['A', 'B', 'B'] and Supplier.objects.count() == 2
    assert ImportExecution.objects.get(pk=new['id']).batch_id == ImportExecution.objects.get(pk=task['id']).batch_id


def test_current_fields_and_data_scope_apply_to_export_and_import(admin_client, admin_user, member_user):
    Supplier.objects.create(code='PRIVATE', name='不可见业务', creator=admin_user)
    Supplier.objects.create(code='OWN', name='自己的业务', creator=member_user)
    fields = {'cg_suppliers': {'read': ['code', 'active'], 'write': ['code', 'active']}}
    client, role = restricted(member_user, fields=fields)
    catalog = client.get(ROOT + 'catalog/?resource=cg_suppliers')
    assert [row['name'] for row in catalog.json()['data']['fields']] == ['code', 'active']
    bad = client.post(ROOT + 'preview/', {'resource': 'cg_suppliers', 'file': upload(), 'options': json.dumps(OPTIONS)}, format='multipart')
    assert bad.status_code == 400
    result = client.get(ROOT + 'export/?resource=cg_suppliers')
    assert result.status_code == 200, result.content
    text = result.content.decode('utf-8-sig')
    assert 'OWN' in text and 'PRIVATE' not in text and 'name' not in text
    assert client.get(ROOT + 'export/?resource=cg_suppliers&name=不可见').status_code == 403
    role.field_permissions = {}; role.save()
    task = execute(client, preview(client, '编号,名称,启用\nPRIVATE,新值,true\n'))
    assert task.result['failed'] == 1 and task.result['created'] == 0
    assert Supplier.objects.get(code='PRIVATE').name == '不可见业务'


def test_permission_revocation_between_preview_submit_and_worker(admin_client, member_user):
    client, role = restricted(member_user)
    plan = preview(client)
    RolePermission.objects.filter(role=role, permission__code='cg_suppliers:update').delete()
    result = client.post(ROOT + f"plans/{plan['id']}/submit/")
    assert result.status_code == 200
    RolePermission.objects.filter(role=role, permission__code='cg_suppliers:create').delete()
    run_data_import(result.json()['data']['id'])
    assert Supplier.objects.count() == 0
    assert TaskRecord.objects.get().status == 'failed'
    assert client.post(ROOT + f"plans/{plan['id']}/submit/").status_code == 403


def test_field_revocation_after_first_committed_row_stops_later_rows(member_user, monkeypatch):
    client, role = restricted(member_user)
    def revoke(**kwargs):
        role.field_permissions = {'cg_suppliers': {'read': ['code', 'name', 'active'], 'write': ['code', 'active']}}
        role.save()
    monkeypatch.setattr(extensions, 'after_save', revoke)
    task = execute(client, preview(client, '编号,名称,启用\nA,甲,true\nB,乙,false\n'))
    assert task.status == 'failed' and task.error == 'execution_not_authorized_committed_rows_preserved'
    assert task.result['created'] == 1 and task.result['processed'] == 1
    assert list(Supplier.objects.values_list('code', flat=True)) == ['A']


def test_source_hash_is_rechecked_by_worker(admin_client):
    task = submit(admin_client, preview(admin_client))
    source = ImportExecution.objects.get(pk=task['id']).batch.source
    with source.file.open('wb') as stream: stream.write(b'code,name\nA,tampered')
    run_data_import(task['id'])
    record = TaskRecord.objects.get(pk=task['id'])
    assert record.status == 'failed' and record.error == 'import_source_changed'
    assert Supplier.objects.count() == 0


@pytest.mark.django_db(transaction=True)
def test_postgres_concurrent_confirmation_registers_one_execution(admin_client, admin_user):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    from django.db import connection, close_old_connections
    from django.contrib.auth import get_user_model
    if connection.vendor != 'postgresql': pytest.skip('确认竞争只在真实PostgreSQL运行')
    plan = preview(admin_client)
    identifier = admin_user.pk; barrier = Barrier(2)
    def confirm(_):
        close_old_connections()
        try:
            client = APIClient(); client.force_authenticate(get_user_model().objects.get(pk=identifier))
            barrier.wait(timeout=15)
            return submit(client, plan)['id']
        finally:
            close_old_connections()
    with ThreadPoolExecutor(2) as pool: results = list(pool.map(confirm, range(2)))
    assert results[0] == results[1] and ImportExecution.objects.count() == 1


def test_personal_mapping_isolation_and_plan_ownership(admin_client, member_user):
    client, _ = restricted(member_user)
    created = admin_client.post(ROOT + 'mappings/', {'resource': 'cg_suppliers', 'name': '财务导入', 'options': OPTIONS}, format='json')
    assert created.status_code == 201, created.content
    identifier = created.json()['data']['id']
    assert client.get(ROOT + 'mappings/?resource=cg_suppliers').json()['data'] == []
    assert client.delete(ROOT + 'mappings/' + identifier + '/').status_code == 404
    plan = preview(admin_client)
    assert client.post(ROOT + f"plans/{plan['id']}/submit/").status_code == 404
    task = execute(admin_client, plan)
    assert client.get(ROOT + f'import/{task.pk}/errors/').status_code == 404
    assert admin_client.delete(ROOT + 'mappings/' + identifier + '/').status_code == 200


def test_expired_or_changed_source_or_schema_is_rejected(admin_client):
    plan = preview(admin_client)
    BusinessImportPlan.objects.filter(pk=plan['id']).update(expires_at=timezone.now() - timedelta(seconds=1))
    assert admin_client.post(ROOT + f"plans/{plan['id']}/submit/").status_code == 400
    plan = preview(admin_client); row = BusinessImportPlan.objects.get(pk=plan['id'])
    with row.source.file.open('wb') as stream:
        stream.write(b'code,name\nchanged,source\n')
    assert admin_client.post(ROOT + f"plans/{plan['id']}/submit/").status_code == 400
    plan = preview(admin_client); row = BusinessImportPlan.objects.get(pk=plan['id'])
    row.options['schema_digest'] = 'changed'; row.save()
    assert admin_client.post(ROOT + f"plans/{plan['id']}/submit/").status_code == 400
    assert TaskRecord.objects.count() == 0


@pytest.mark.parametrize('invalid', [dict(OPTIONS, key=[]), dict(OPTIONS, mode='overwrite'), dict(OPTIONS, mapping={'编号': 'id'}), dict(OPTIONS, mapping={'编号': 'code'}), dict(OPTIONS, mapping={'编号': 'code', '名称': 'code'})])
def test_invalid_mapping_never_creates_plan(admin_client, invalid):
    response = admin_client.post(ROOT + 'preview/', {'resource': 'cg_suppliers', 'file': upload(), 'options': json.dumps(invalid)}, format='multipart')
    assert response.status_code == 400, response.content
    assert BusinessImportPlan.objects.count() == 0


@pytest.mark.parametrize('raw', [b'code,code\nA,A', b'code\n', b'code,name\nA,B,C', b'code\n' + b'A\n' * 5001, b'code\nA\0'])
def test_invalid_csv_rejected(raw):
    with pytest.raises(ValidationError):
        parse_rows(raw, 'rows.csv')


def test_xlsx_formula_and_multiple_sheets_rejected_and_typed_values_supported():
    book = Workbook(); book.active.append(['code', 'active']); book.active.append(['A', True])
    buffer = io.BytesIO(); book.save(buffer)
    assert parse_rows(buffer.getvalue(), 'rows.xlsx')[1] == [{'code': 'A', 'active': 'True'}]
    book.active.cell(2, 1, '=1+1'); buffer = io.BytesIO(); book.save(buffer)
    with pytest.raises(ValidationError): parse_rows(buffer.getvalue(), 'rows.xlsx')
    book.active.cell(2, 1, 'A'); book.create_sheet(); buffer = io.BytesIO(); book.save(buffer)
    with pytest.raises(ValidationError): parse_rows(buffer.getvalue(), 'rows.xlsx')


def test_exports_match_filters_escape_formulas_and_enforce_limits(admin_client, admin_user, monkeypatch):
    from apps.data_io import business
    Supplier.objects.create(code='A', name='=HYPERLINK("bad")', creator=admin_user)
    Supplier.objects.create(code='B', name='普通业务', creator=admin_user)
    response = admin_client.get(ROOT + 'export/?resource=cg_suppliers&code=A')
    assert response.status_code == 200, response.content
    rows = list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))
    assert len(rows) == 2 and rows[1][1].startswith("'=HYPERLINK")
    response = admin_client.get(ROOT + 'export/?resource=cg_suppliers&code=A&export_format=xlsx')
    assert response.status_code == 200, response.content
    book = load_workbook(io.BytesIO(response.content)); assert book.active.cell(2, 2).data_type != 'f'
    book.close()
    monkeypatch.setattr(business, 'MAX_ROWS', 1)
    assert admin_client.get(ROOT + 'export/?resource=cg_suppliers').status_code == 400
    for extra in ['&code=A&code=B', '&page=1', '&nonsense=1', '&export_format=pdf']:
        assert admin_client.get(ROOT + 'export/?resource=cg_suppliers' + extra).status_code == 400


def test_tenant_plan_mapping_and_records_do_not_cross(admin_client, admin_user, settings):
    settings.ADMINCORE_MULTI_TENANT_ENABLED = True
    one, two = Tenant.objects.create(name='一', slug='io-one'), Tenant.objects.create(name='二', slug='io-two')
    # 直接调用服务避免租户中间件再次选择；所有存储均使用真实租户管理器。
    from apps.data_io.business import preflight, options_for
    with tenant_context(one.pk):
        Supplier.objects.create(code='A', name='租户一', creator=admin_user)
        ImportMapping.objects.create(creator=admin_user, resource='cg_suppliers', name='租户一模板', options=OPTIONS)
    with tenant_context(two.pk):
        assert ImportMapping.objects.count() == 0
        options = options_for(admin_user, 'cg_suppliers', OPTIONS, ['编号', '名称', '启用'])
        _, result = preflight(admin_user, 'cg_suppliers', [{'编号': 'A', '名称': '租户二', '启用': 'true'}], options)
        assert result['created'] == 1 and Supplier.objects.count() == 0


def test_expired_preview_metadata_uses_existing_retention(admin_client):
    from apps.tasks.lifecycle import clean
    plan = preview(admin_client)
    BusinessImportPlan.objects.filter(pk=plan['id']).update(expires_at=timezone.now() - timedelta(days=3650))
    dry = clean(dry_run=True); assert dry['eligible']['import_plans'] == 1
    assert BusinessImportPlan.objects.count() == 1
    result = clean(); assert result['deleted']['import_plans'] == 1 and BusinessImportPlan.objects.count() == 0


def test_unregistered_or_complex_resources_cannot_use_single_table_endpoints(admin_client):
    for resource in ['users', 'cg_categories', 'unknown']:
        assert admin_client.get(ROOT + 'catalog/?resource=' + resource).status_code == 404
    assert admin_client.post('/api/v1/data-io/import/', {'resource': 'cg_suppliers', 'file': upload()}, format='multipart').status_code == 400
