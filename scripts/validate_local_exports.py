"""真实 PostgreSQL/Redis/Nginx 当前页导出验收；独立数据库，不读取正式管理员凭据。"""
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
import hashlib
import json
import os
import secrets
import shutil
import socket
import subprocess
import uuid

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PROJECT = 'admincore-local-exports-' + uuid.uuid4().hex[:8]
FOLDER = ROOT / '.codex-validation' / PROJECT
PORTS = [18470, 18471]
BASE = 'http://127.0.0.1:18471/api/v1/'
CHECKS = []


def run(arguments, *, input=None, allowed_failure=False):
    result = subprocess.run(arguments, cwd=ROOT, capture_output=True, text=True, input=input,
                            encoding='utf-8', errors='replace', timeout=360)
    with (FOLDER / 'contract.log').open('a', encoding='utf-8') as log:
        log.write(result.stdout + result.stderr)
    if result.returncode and not allowed_failure:
        raise RuntimeError('本地导出隔离验收失败，见 contract.log')
    return result


def compose(mode, *arguments, input=None, allowed_failure=False):
    command = ['docker', 'compose', '-p', PROJECT, '--env-file', str(FOLDER / 'fixture.env')]
    for file in ['docker-compose.yml', 'compose.minimal.yml', str(FOLDER / 'override.yml')]:
        command += ['-f', file]
    if mode == 'off': command += ['-f', str(FOLDER / 'off.yml')]
    return run(command + list(arguments), input=input, allowed_failure=allowed_failure)


def check(label, condition=True):
    if not condition: raise AssertionError(label)
    CHECKS.append(label); print('通过：' + label, flush=True)


def snippet(code, mode='on'):
    prefix = "import os; os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings'); import django; django.setup(); from django.conf import settings; assert settings.DATABASES['default']['NAME']=='local_exports_fixture'\n"
    return compose(mode, 'exec', '-T', 'backend', 'python', '-c', 'import sys; exec(sys.stdin.read())', input=prefix + code).stdout.strip()


def api(path, token=None, method='GET', data=None, expected=200):
    response = requests.request(method, BASE + path, headers={'Authorization': 'Bearer ' + token} if token else {}, json=data, timeout=30)
    check(method + ' ' + path + ' → ' + str(expected), response.status_code == expected)
    return response


def inspect_downloads():
    results = []
    for path in sorted(FOLDER.glob('*.xlsx')):
        workbook = load_workbook(path, data_only=False, read_only=False, keep_links=True)
        assert workbook.sheetnames == ['Data', 'Export info']
        values = []
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    assert cell.data_type != 'f' and not cell.hyperlink
                    values.append(str(cell.value))
        assert 'member-private@example.test' not in values and '13912345678' not in values
        with ZipFile(path) as archive:
            assert not any(name.lower().endswith('vbaproject.bin') for name in archive.namelist())
            for name in archive.namelist():
                if name.endswith('.rels'):
                    assert all(row.attrib.get('TargetMode') != 'External' for row in ET.fromstring(archive.read(name)))
        results.append({'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                        'rows': workbook['Data'].max_row - 1, 'columns': workbook['Data'].max_column})
    check('独立 openpyxl/ZIP 验证16个真实HTTP列表下载，无公式/外链/宏/原始联系方式', len(results) == 16)
    (FOLDER / 'workbook-report.json').write_text(json.dumps({'passed': True, 'files': results}, indent=2), encoding='utf-8')


def main():
    for port in PORTS:
        with socket.socket() as listener: listener.bind(('127.0.0.1', port))
    FOLDER.mkdir(parents=True)
    password, changed = secrets.token_urlsafe(40), secrets.token_urlsafe(40)
    env_file = FOLDER / 'fixture.env'
    environment = {'DJANGO_SECRET_KEY': secrets.token_urlsafe(60), 'DB_PASSWORD': secrets.token_urlsafe(40),
        'DB_NAME': 'local_exports_fixture', 'DB_USER': 'postgres', 'ADMINCORE_ADMIN_PASSWORD': password,
        'DJANGO_ALLOWED_HOSTS': 'localhost,127.0.0.1,backend', 'DJANGO_ENV': 'production', 'DJANGO_DEBUG': 'False',
        'BACKEND_PORT': '18470', 'FRONTEND_PORT': '18471', 'AUTH_LOGIN_CAPTCHA_ENABLED': 'False', 'USE_TLS': 'False'}
    env_file.write_text(''.join(key + '=' + value + '\n' for key, value in environment.items()), encoding='utf-8')
    (FOLDER / 'override.yml').write_text(f'''services:
  redis:
    profiles: !reset []
  backend:
    image: {PROJECT}-backend
    env_file: !override
      - "{env_file.as_posix()}"
    environment:
      ADMINCORE_MODULES: files,tasks,data_io
    depends_on: !override
      db:
        condition: service_healthy
      redis:
        condition: service_healthy
  frontend:
    image: {PROJECT}-frontend
    ports: !override
      - "127.0.0.1:18471:80"
''', encoding='utf-8')
    (FOLDER / 'off.yml').write_text('services:\n  backend:\n    environment:\n      ADMINCORE_MODULES: files,tasks\n', encoding='utf-8')
    report = {'passed': False, 'project': PROJECT, 'ports': PORTS}
    try:
        compose('on', 'build', 'backend', 'frontend')
        compose('on', 'up', '-d', '--wait', '--wait-timeout', '150', 'db', 'redis', 'backend', 'frontend')
        check('四服务验收不启动导出worker/beat/外部转换器', set(compose('on', 'ps', '--status', 'running', '--services').stdout.split()) == {'db', 'redis', 'backend', 'frontend'})
        token = api('auth/login/', method='POST', data={'username': 'admin', 'password': password}).json()['access']
        api('auth/change_password/', token, 'POST', {'old_password': password, 'new_password': changed, 'new_password2': changed})
        api('auth/login/', method='POST', data={'username': 'admin', 'password': changed})
        retained_code = "from django.contrib.auth import get_user_model; from apps.files.models import FileModel; import hashlib,json; u=get_user_model().objects.get(username='local_export_member'); f=FileModel.objects.get(original_name='export-fixture.txt'); print(hashlib.sha256(json.dumps([str(u.pk),u.email,u.phone,str(f.pk),f.lifecycle_version,f.file.read().decode()],sort_keys=True).encode()).hexdigest())"
        snippet("from django.contrib.auth import get_user_model; from apps.files.models import FileModel; from django.core.files.base import ContentFile; u=get_user_model().objects.create_user(username='local_export_member',email='member-private@example.test',phone='13912345678',password=None); FileModel.objects.create(original_name='export-fixture.txt',file_type='txt',size=7,md5='0'*32,uploader=get_user_model().objects.get(username='admin'),file=ContentFile(b'fixture',name='export-fixture.txt'))")
        original = snippet(retained_code)
        (FOLDER / 'browser-access.json').write_text(json.dumps({'password': changed}), encoding='utf-8')
        node_bin = ROOT / '.codex-validation/toolchains/node-v24.21.0-win-x64/node.exe'
        node = str(node_bin) if os.name == 'nt' and node_bin.exists() else shutil.which('node')
        if not node: raise RuntimeError('真实浏览器验收需要Node.js')
        def browser(mode):
            run([node, str(ROOT / 'scripts/local_export_browser.mjs'), str(ROOT / 'frontend'), str(FOLDER), mode])
        browser('on'); check('桌面/手机真实登录、四列表私有字段与本地下载')
        compose('off', 'up', '-d', '--wait', '--wait-timeout', '150', 'backend')
        browser('off'); check('关闭data_io撤下本地导出入口，原列表与数据保留')
        check('关闭期间原账号/文件内容与版本一致', snippet(retained_code, 'off') == original)
        compose('on', 'up', '-d', '--wait', '--wait-timeout', '150', 'backend')
        browser('reenabled'); check('重新启用后四列表可导出，原数据保持一致', snippet(retained_code) == original)
        inspect_downloads()
        report['passed'] = True
    finally:
        report['checks'] = CHECKS
        if report['passed']:
            compose('on', 'down', '-v', '--remove-orphans')
            for name in ['fixture.env', 'browser-access.json']: (FOLDER / name).unlink(missing_ok=True)
            report['cleaned'] = ['isolated containers and volumes', 'fixture credentials']
        else:
            compose('on', 'logs', '--no-color', allowed_failure=True); compose('on', 'stop', allowed_failure=True)
        (FOLDER / 'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        print(json.dumps({'evidence': str(FOLDER), **report}, ensure_ascii=False), flush=True)


if __name__ == '__main__': main()
